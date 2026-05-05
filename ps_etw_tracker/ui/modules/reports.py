import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from ui.theme import C, F
from pathlib import Path
from datetime import date
import os
import sys

EXPORT_DIR = Path(__file__).resolve().parent.parent.parent / "exports"
EXPORT_DIR.mkdir(parents=True, exist_ok=True)


class ReportsModule:
    def __init__(self, parent, app):
        self.parent = parent
        self.app = app

    def render(self):
        tk.Label(self.parent, text="Generate Reports",
                 font=F["h2"], bg=C["bg"]).pack(anchor="w", pady=(0, 16))

        REPORTS = [
            ("📊", "Project Summary (Excel)",
             "Multi-sheet: KPIs, Activities, Materials, Delays",
             lambda: export_summary(self.app)),
            ("📋", "Activities Export (Excel)",
             "All activities with status, dates, hours",
             lambda: export_activities(self.app)),
            ("🔩", "Materials Report (Excel)",
             "Procurement register with risk flags and alerts",
             lambda: export_materials(self.app)),
            ("⚠",  "Delay Log (Excel)",
             "All delays, root causes, action owners, resolutions",
             lambda: export_delays(self.app)),
            ("📅", "Gantt Chart (PostScript)",
             "Export current Gantt tab — Export PS button",
             lambda: messagebox.showinfo("Info",
                 "Open the Gantt tab and use Export PS.")),
        ]

        grid = tk.Frame(self.parent, bg=C["bg"])
        grid.pack(fill="both")

        for i, (ico, title, desc, action) in enumerate(REPORTS):
            row, col = divmod(i, 2)
            card = tk.Frame(grid, bg=C["card"], padx=14, pady=14,
                            cursor="hand2")
            card.grid(row=row, column=col, sticky="nsew",
                      padx=(0 if col == 0 else 8, 0),
                      pady=(0 if row == 0 else 8, 0))
            card.configure(highlightbackground=C["border"],
                           highlightthickness=1)
            grid.grid_columnconfigure(col, weight=1)

            ico_frame = tk.Frame(card, bg=C["red_light"], width=44, height=44)
            ico_frame.pack(side="left", padx=(0, 12))
            ico_frame.pack_propagate(False)
            tk.Label(ico_frame, text=ico, font=(F["h1"][0], 18),
                     bg=C["red_light"]).pack(expand=True)

            text_f = tk.Frame(card, bg=C["card"])
            text_f.pack(side="left", fill="both", expand=True)
            tk.Label(text_f, text=title, font=F["h3"],
                     bg=C["card"], fg="#111111").pack(anchor="w")
            tk.Label(text_f, text=desc, font=F["small"],
                     bg=C["card"], fg=C["muted"],
                     wraplength=250, justify="left").pack(anchor="w", pady=(2, 0))

            the_action = action
            card.bind("<Button-1>", lambda e, a=the_action: a())
            card.bind("<Enter>",
                lambda e, c=card: c.configure(
                    highlightbackground=C["red"], highlightthickness=2))
            card.bind("<Leave>",
                lambda e, c=card: c.configure(
                    highlightbackground=C["border"], highlightthickness=1))

        info = tk.Frame(self.parent, bg="#FAFAFA", padx=14, pady=12)
        info.pack(fill="x", pady=(16, 0))
        info.configure(highlightbackground=C["border"], highlightthickness=1)
        tk.Label(info, text=f"📁  Exports saved to:  {EXPORT_DIR}",
                 font=F["small"], bg="#FAFAFA", fg=C["text"]).pack(anchor="w")
        ttk.Button(info, text="Open Folder", style="Ghost.TButton",
                   command=lambda: _open_folder(EXPORT_DIR)).pack(
            anchor="w", pady=(4, 0))


def export_activities(app) -> str | None:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        messagebox.showerror("Missing Library",
            "openpyxl is required.\nRun: pip install openpyxl")
        return None

    proj_id = app.active_project_id
    if not proj_id:
        return None

    proj = next((p for p in app.projects if p["id"] == proj_id), {})
    acts = app.activities

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Activities"
    ws.freeze_panes = "D4"

    ws.merge_cells("A1:T1")
    ws["A1"] = f"PS-ETW Engine Build-Up Tracker — {proj.get('code','')} | {proj.get('name','')}"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="ED0007")
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:T2")
    ws["A2"] = (f"Generated: {date.today().strftime('%d-%b-%Y')} | "
                f"Target: {proj.get('target_finish','—')}")
    ws["A2"].font = Font(name="Calibri", size=9)
    ws["A2"].fill = PatternFill("solid", fgColor="F5F5F5")

    HEADERS = ["#", "Major Phase", "Phase", "Activity", "Sub-Activity", "Tech",
               "Dept", "Priority", "Plan Start", "Plan Finish", "Dur(d)", "Float(d)",
               "% Done", "Status", "Est.h", "Act.h", "Act.Start", "Act.Finish",
               "Delay(d)", "Alert"]
    WIDTHS = [5, 18, 18, 42, 20, 16, 8, 10, 12, 12, 8, 8, 9, 14, 7, 7, 12, 12, 8, 8]

    thin = Side(border_style="thin", color="E8E8E8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="37474F")
    hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col, (h, w) in enumerate(zip(HEADERS, WIDTHS), 1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = hdr_align
        cell.border = border
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[3].height = 22

    STATUS_FILLS = {
        "Complete":    PatternFill("solid", fgColor="E8F5E9"),
        "In Progress": PatternFill("solid", fgColor="E3F2FD"),
        "Blocked":     PatternFill("solid", fgColor="FFEBEE"),
        "Delayed":     PatternFill("solid", fgColor="FFF3E0"),
        "On Hold":     PatternFill("solid", fgColor="F3E5F5"),
    }
    std_font = Font(name="Calibri", size=10)

    cur_major = None
    row = 4
    for a in acts:
        if a["major_phase"] != cur_major:
            cur_major = a["major_phase"]
            ws.merge_cells(f"A{row}:T{row}")
            ws.cell(row=row, column=1, value=f"⬛  {cur_major}")
            ws.cell(row=row, column=1).font = Font(
                name="Calibri", size=10, bold=True, color="FFFFFF")
            ws.cell(row=row, column=1).fill = PatternFill(
                "solid", fgColor="37474F")
            ws.row_dimensions[row].height = 18
            row += 1

        fill = STATUS_FILLS.get(a["status"], PatternFill("solid", fgColor="FFFFFF"))
        row_data = [
            a["seq_num"], a["major_phase"], a["phase"], a["activity_name"],
            a.get("sub_activity", ""), a.get("technician", ""), a["dept"], a["priority"],
            a.get("plan_start", ""), a.get("plan_finish", ""),
            a.get("duration_d", 0), a.get("float_d", 0),
            f"{a.get('pct_complete', 0)}%", a["status"],
            a.get("est_hours", 0), a.get("actual_hours", 0),
            a.get("actual_start", ""), a.get("actual_finish", ""),
            a.get("delay_d", 0), a.get("alert_level", "OK"),
        ]
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.font = std_font
            cell.fill = fill
            cell.border = border
            cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[row].height = 18
        row += 1

    fname = EXPORT_DIR / f"{proj_id}_Activities_{date.today().isoformat()}.xlsx"
    wb.save(fname)
    _open_file_or_folder(fname)
    app.toast.show("Exported", f"Activities saved: {fname.name}", "ok")
    return str(fname)


def export_materials(app) -> str | None:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Border, Side
    except ImportError:
        messagebox.showerror("Missing Library", "Install openpyxl: pip install openpyxl")
        return None

    proj_id = app.active_project_id
    if not proj_id:
        return None
    proj = next((p for p in app.projects if p["id"] == proj_id), {})
    mats = app.materials

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Materials"
    ws.merge_cells("A1:S1")
    ws["A1"] = f"PS-ETW Materials — {proj.get('code','')} | {proj.get('name','')}"
    ws["A1"].font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="ED0007")

    HEADERS = ["#", "Part Number", "Description", "Linked Act.", "Qty", "Owner",
               "Crit.", "Need-By", "Status", "Supplier", "Tracking",
               "Promised", "Received", "Risk", "Lead(d)", "Value(€)",
               "Notes", "Latest Order", "Order Alert"]
    thin = Side(border_style="thin", color="E8E8E8")
    bdr = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=2, column=col, value=h)
        cell.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="37474F")
        cell.border = bdr

    STATUS_FILLS = {
        "Not Ordered": PatternFill("solid", fgColor="FFEBEE"),
        "In Transit":  PatternFill("solid", fgColor="E3F2FD"),
        "Available":   PatternFill("solid", fgColor="E8F5E9"),
        "Arrived":     PatternFill("solid", fgColor="FFF3E0"),
    }
    for r, m in enumerate(mats, 3):
        fill = STATUS_FILLS.get(m["status"], PatternFill("solid", fgColor="FFFFFF"))
        row_data = [
            m["id"], m["part_number"], m["description"],
            m.get("linked_activity_id", ""), m["quantity"], m["ownership"],
            m["criticality"], m["need_by_date"], m["status"], m["supplier"],
            m["tracking_ref"], m["promised_date"], m["received_date"],
            m["risk_flag"], m["lead_time_d"], m["value_eur"],
            m["notes"], m["latest_order_date"], m["order_alert"]
        ]
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=r, column=col, value=val)
            cell.fill = fill
            cell.font = Font(name="Calibri", size=10)
            cell.border = bdr

    fname = EXPORT_DIR / f"{proj_id}_Materials_{date.today().isoformat()}.xlsx"
    wb.save(fname)
    _open_file_or_folder(fname)
    app.toast.show("Exported", f"Materials saved: {fname.name}", "ok")
    return str(fname)


def export_delays(app) -> str | None:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Border, Side
    except ImportError:
        messagebox.showerror("Missing Library", "Install openpyxl: pip install openpyxl")
        return None

    proj_id = app.active_project_id
    if not proj_id:
        return None
    proj = next((p for p in app.projects if p["id"] == proj_id), {})
    delays = app.delays

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Delays & Actions"
    ws.merge_cells("A1:M1")
    ws["A1"] = f"PS-ETW Delays & Actions — {proj.get('code','')} | {proj.get('name','')}"
    ws["A1"].font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="ED0007")

    HEADERS = ["#", "Activity", "Phase", "Tech", "Action Required", "Root Cause",
               "Action Owner", "Target Date", "Resolution", "Status", "Delay(hrs)",
               "Created", "Updated"]
    thin = Side(border_style="thin", color="E8E8E8")
    bdr = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=2, column=col, value=h)
        cell.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="37474F")
        cell.border = bdr

    for r, d in enumerate(delays, 3):
        row_data = [
            d.get("seq_num", ""), d.get("activity_name", ""), d.get("phase", ""),
            d.get("technician", ""), d.get("action_required", ""),
            d.get("root_cause", ""), d.get("action_owner", ""),
            d.get("target_date", ""), d.get("resolution", ""),
            d.get("delay_status", ""), d.get("total_delay_hrs", 0),
            d.get("created_at", ""), d.get("updated_at", ""),
        ]
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=r, column=col, value=val)
            cell.font = Font(name="Calibri", size=10)
            cell.border = bdr

    fname = EXPORT_DIR / f"{proj_id}_Delays_{date.today().isoformat()}.xlsx"
    wb.save(fname)
    _open_file_or_folder(fname)
    app.toast.show("Exported", f"Delays saved: {fname.name}", "ok")
    return str(fname)


def export_summary(app) -> str | None:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        messagebox.showerror("Missing", "Install openpyxl")
        return None

    proj_id = app.active_project_id
    if not proj_id:
        return None
    proj = next((p for p in app.projects if p["id"] == proj_id), {})
    acts = app.activities
    mats = app.materials
    dels = app.delays
    total = len(acts)
    done = sum(1 for a in acts if a["status"] == "Complete")

    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "PS-ETW ENGINE BUILD-UP TRACKER"
    ws["A1"].font = Font(name="Calibri", size=18, bold=True, color="ED0007")
    ws["A3"] = f"Project Code:   {proj.get('code','')}"
    ws["A4"] = f"Project Name:   {proj.get('name','')}"
    ws["A5"] = f"Customer:       {proj.get('customer','')}"
    ws["A6"] = f"Engine:         {proj.get('engine_type','')}"
    ws["A7"] = f"PM:             {proj.get('pm_name','')}"
    ws["A8"] = f"Start Date:     {proj.get('start_date','')}"
    ws["A9"] = f"Target Finish:  {proj.get('target_finish','')}"
    ws["A10"] = f"Report Date:    {date.today().strftime('%d-%b-%Y')}"
    ws["A12"] = (
        f"Progress: {done}/{total} activities complete "
        f"({round(done/total*100) if total else 0}%)")
    ws["A12"].font = Font(name="Calibri", size=12, bold=True)

    ws2 = wb.create_sheet("Activities")
    hdr = ["#", "Phase", "Activity", "Status", "Plan Start", "Plan Finish", "%"]
    ws2.append(hdr)
    for a in acts:
        ws2.append([
            a.get("seq_num"), a.get("phase"), a.get("activity_name"),
            a.get("status"), a.get("plan_start"), a.get("plan_finish"),
            a.get("pct_complete"),
        ])

    ws3 = wb.create_sheet("Materials")
    ws3.append(["#", "Part", "Description", "Status", "Risk"])
    for m in mats:
        ws3.append([
            m.get("id"), m.get("part_number"), m.get("description"),
            m.get("status"), m.get("risk_flag"),
        ])

    ws4 = wb.create_sheet("Delays")
    ws4.append(["Activity", "Phase", "Status", "Root Cause", "Action"])
    for d in dels:
        ws4.append([
            d.get("activity_name"), d.get("phase"), d.get("delay_status"),
            d.get("root_cause"), d.get("action_required"),
        ])

    fname = EXPORT_DIR / f"{proj_id}_Summary_{date.today().isoformat()}.xlsx"
    wb.save(fname)
    _open_file_or_folder(fname)
    app.toast.show("Summary Exported", fname.name, "ok")
    return str(fname)


def _open_file_or_folder(path: Path):
    try:
        if sys.platform == "win32":
            os.startfile(str(path))
        elif sys.platform == "darwin":
            import subprocess
            subprocess.Popen(["open", str(path)])
        else:
            import subprocess
            subprocess.Popen(["xdg-open", str(path)])
    except Exception:
        pass


def _open_folder(path: Path):
    _open_file_or_folder(path)
